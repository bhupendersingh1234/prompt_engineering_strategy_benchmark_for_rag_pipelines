"""
Generates the First Review PPT (10 slides, ~50% implementation checkpoint)
for "Prompt Engineering Strategy Benchmark for RAG Pipelines".

Deliberately does NOT include the full 4-strategy comparison / final
numbers from README.md Results section -- that's reserved for the final
review. This deck shows: problem, architecture, methodology, what's been
built and validated so far (pipeline + evaluation framework + judge
calibration), and what's planned next.

Usage:
    python presentation/generate_first_review_ppt.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

OUT_PATH = Path(__file__).resolve().parent / "First_Review_Prompt_Engineering_RAG.pptx"

# ---------------------------------------------------------------------------
# Palette / type system (mirrors the zeroth-review title-slide reference)
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x0D, 0x3B, 0x66)
NAVY_DARK = RGBColor(0x08, 0x28, 0x47)
INK = RGBColor(0x1A, 0x1D, 0x22)
INK_SOFT = RGBColor(0x4A, 0x4F, 0x58)
MUTED = RGBColor(0x8A, 0x8F, 0x98)
LINE = RGBColor(0xD8, 0xDC, 0xE2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAGE_BG = RGBColor(0xFF, 0xFF, 0xFF)
PINK_BG = RGBColor(0xFB, 0xE4, 0xE4)
PINK_INK = RGBColor(0xA9, 0x33, 0x3F)
ACCENT_GOLD = RGBColor(0xC7, 0x9A, 0x3D)
DONE_GREEN = RGBColor(0x1E, 0x7A, 0x4C)
DONE_GREEN_BG = RGBColor(0xE3, 0xF3, 0xE9)
PLAN_AMBER = RGBColor(0x9A, 0x6A, 0x12)
PLAN_AMBER_BG = RGBColor(0xFB, 0xF0, 0xDE)

SERIF = "Georgia"
SANS = "Segoe UI"
MONO = "Consolas"

REVIEW_LABEL = "First Review"
FOOTER_TEXT = f"B.Tech Project – I | {REVIEW_LABEL} | School of Computer Science and Engineering | VIT Chennai"

TEAM = [
    ("23BAI1057", "Astitva Thakkar"),
    ("23BAI1528", "Parth Sharma"),
    ("23BRS1107", "Bhupender Singh"),
]
FACULTY = {
    "Name": "Manmohan Sharma",
    "ERP ID": "53622",
    "School": "SENSE (School of Electronics and Communication Engineering)",
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAGE_BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    # send to back
    spTree = slide.shapes._spTree
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)
    return slide


def add_textbox(slide, left, top, width, height, text, size=14, color=INK,
                 bold=False, italic=False, font=SANS, align=PP_ALIGN.LEFT,
                 anchor=MSO_ANCHOR.TOP, line_spacing=1.0, wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return box


def add_bullets(slide, left, top, width, height, items, size=15, color=INK,
                 font=SANS, gap_pt=10, marker="—", marker_color=None,
                 bold_lead=None, line_spacing=1.12):
    """items: list of str, or (lead_bold, rest) tuples."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(gap_pt)
        r0 = p.add_run()
        r0.text = f"{marker}  "
        r0.font.size = Pt(size)
        r0.font.name = font
        r0.font.color.rgb = marker_color or color
        r0.font.bold = True
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run()
            r1.text = lead
            r1.font.size = Pt(size)
            r1.font.name = font
            r1.font.bold = True
            r1.font.color.rgb = color
            if rest:
                r2 = p.add_run()
                r2.text = rest
                r2.font.size = Pt(size)
                r2.font.name = font
                r2.font.color.rgb = color
        else:
            r1 = p.add_run()
            r1.text = item
            r1.font.size = Pt(size)
            r1.font.name = font
            r1.font.color.rgb = color
    return box


def add_rect(slide, left, top, width, height, fill=None, line_color=None,
             line_width=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    if radius:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is not None:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    if line_color is not None:
        shp.line.color.rgb = line_color
        shp.line.width = line_width or Pt(0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_line(slide, left, top, width, color=LINE, weight=Pt(1)):
    ln = slide.shapes.add_connector(1, left, top, left + width, top)
    ln.line.color.rgb = color
    ln.line.width = weight
    return ln


def add_footer(slide, page_no):
    bar = add_rect(slide, Inches(0.5), Inches(6.95), Inches(12.33), Inches(0.4), fill=PINK_BG)
    tf = bar.text_frame
    tf.margin_left = Inches(0.18)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = f"⚑  {FOOTER_TEXT}"
    r.font.size = Pt(10.5)
    r.font.name = SERIF
    r.font.color.rgb = PINK_INK
    # page number, right aligned
    pg = add_textbox(slide, Inches(12.0), Inches(6.98), Inches(0.7), Inches(0.32),
                      f"{page_no:02d}", size=10.5, color=PINK_INK, font=MONO,
                      align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    return bar


def add_header(slide, eyebrow, title, subtitle=None):
    add_textbox(slide, Inches(0.6), Inches(0.35), Inches(11.5), Inches(0.32),
                eyebrow.upper(), size=11.5, color=NAVY, bold=True, font=SANS)
    add_textbox(slide, Inches(0.6), Inches(0.68), Inches(11.8), Inches(0.75),
                title, size=28, color=INK, bold=False, font=SERIF)
    y = Inches(1.42)
    if subtitle:
        add_textbox(slide, Inches(0.6), y, Inches(11.8), Inches(0.5),
                    subtitle, size=13.5, color=INK_SOFT, font=SANS, line_spacing=1.2)
        y = Inches(1.85)
    add_line(slide, Inches(0.6), Inches(1.34) if not subtitle else Inches(1.78),
             Inches(12.1), color=LINE, weight=Pt(1))
    return y


def status_pill(slide, left, top, label, kind="done"):
    fill = DONE_GREEN_BG if kind == "done" else PLAN_AMBER_BG
    ink = DONE_GREEN if kind == "done" else PLAN_AMBER
    w = Inches(0.24 + 0.092 * len(label))
    pill = add_rect(slide, left, top, w, Inches(0.30), fill=fill, radius=0.5)
    tf = pill.text_frame
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.1)
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = ("✓ DONE" if kind == "done" else "○ PLANNED")
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.name = SANS
    r.font.color.rgb = ink
    return w


# ===========================================================================
# SLIDE 1 -- Title (mirrors the zeroth-review reference image)
# ===========================================================================
s = add_slide()
add_textbox(s, Inches(0), Inches(0.75), SW, Inches(0.4),
            f"VIT CHENNAI  —  B.Tech Project – I  —  {REVIEW_LABEL}",
            size=15, color=INK_SOFT, font=SERIF, align=PP_ALIGN.CENTER)
add_textbox(s, Inches(0.8), Inches(1.15), Inches(11.73), Inches(1.0),
            "Prompt Engineering Strategy Benchmark for RAG Pipelines",
            size=34, color=INK, font=SERIF, align=PP_ALIGN.CENTER)
add_line(s, Inches(0.45), Inches(2.35), Inches(12.43), color=LINE, weight=Pt(1))

team_box = add_rect(s, Inches(0.45), Inches(2.68), Inches(5.75), Inches(2.15), fill=NAVY)
tf = team_box.text_frame
tf.margin_left = Inches(0.28)
tf.margin_top = Inches(0.22)
tf.margin_right = Inches(0.2)
p = tf.paragraphs[0]
r = p.add_run()
r.text = "Project Team"
r.font.size = Pt(16)
r.font.name = SERIF
r.font.color.rgb = WHITE
for reg, name in TEAM:
    p = tf.add_paragraph()
    p.space_before = Pt(14)
    r1 = p.add_run()
    r1.text = f"{reg}"
    r1.font.bold = True
    r1.font.size = Pt(13.5)
    r1.font.name = SANS
    r1.font.color.rgb = WHITE
    r2 = p.add_run()
    r2.text = f"  —  {name}"
    r2.font.size = Pt(13.5)
    r2.font.name = SANS
    r2.font.color.rgb = RGBColor(0xDD, 0xE6, 0xF0)

fac_left = Inches(6.5)
add_textbox(s, fac_left, Inches(2.68), Inches(5.3), Inches(0.4),
            "Faculty Guide", size=16, color=INK, font=SERIF)
fy = 3.18
for k, v in FACULTY.items():
    add_textbox(s, fac_left, Inches(fy), Inches(5.3), Inches(0.35),
                f"{k}: {v}", size=13, color=INK_SOFT, font=SANS)
    fy += 0.4
add_textbox(s, fac_left, Inches(fy + 0.12), Inches(5.3), Inches(0.35),
            "Signature: ________________________", size=13, color=INK_SOFT, font=SANS)

footer1 = add_rect(s, Inches(0.45), Inches(6.55), Inches(12.43), Inches(0.55), fill=PINK_BG)
tf = footer1.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.22)
p = tf.paragraphs[0]
r = p.add_run()
r.text = f"⚑  {FOOTER_TEXT}"
r.font.size = Pt(12.5)
r.font.name = SERIF
r.font.color.rgb = PINK_INK


# ===========================================================================
# SLIDE 2 -- Problem Statement & Motivation
# ===========================================================================
s = add_slide()
add_header(s, "01 · Problem", "Retrieval fixes hallucination — prompting still decides how well",
           "RAG grounds an LLM in real documents, but no controlled study isolates the prompt as the only variable.")
add_bullets(s, Inches(0.6), Inches(2.15), Inches(6.9), Inches(4.0), [
    ("Large Language Models ", "hallucinate when answering from parameters alone — RAG mitigates this by retrieving relevant context before generation."),
    ("Answer quality still depends heavily ", "on how that retrieved context is handed to the model — the prompting strategy."),
    ("Published comparisons rarely isolate the prompt", " — they change retrieval, embeddings, or the base model at the same time, so the reported gain can't be attributed to the prompt alone."),
    ("This project's thesis: ", "hold chunking, embeddings, vector store, top-k, and the base LLM strictly constant, and vary only the prompting strategy."),
], size=15.5)
card = add_rect(s, Inches(7.85), Inches(2.15), Inches(4.9), Inches(3.9), fill=RGBColor(0xF5, 0xF7, 0xFA), line_color=LINE)
tf = card.text_frame
tf.margin_left = Inches(0.32)
tf.margin_top = Inches(0.28)
tf.margin_right = Inches(0.28)
p = tf.paragraphs[0]
r = p.add_run(); r.text = "Four strategies under test"; r.font.bold = True; r.font.size = Pt(14); r.font.name = SANS; r.font.color.rgb = NAVY
for label in ["Zero-shot", "Few-shot", "Chain-of-Thought", "Structured Output"]:
    p = tf.add_paragraph(); p.space_before = Pt(10)
    r = p.add_run(); r.text = f"•  {label}"; r.font.size = Pt(14); r.font.name = SANS; r.font.color.rgb = INK
p = tf.add_paragraph(); p.space_before = Pt(16)
r = p.add_run(); r.text = "Dataset: RAGTruth (wandb/RAGTruth-processed, HF)"; r.font.size = Pt(12.5); r.font.italic = True; r.font.name = SANS; r.font.color.rgb = INK_SOFT
add_footer(s, 2)


# ===========================================================================
# SLIDE 3 -- Objectives & Current Status
# ===========================================================================
s = add_slide()
add_header(s, "02 · Objectives", "Four objectives — where each one stands today")

objs = [
    ("Build a common, reusable RAG pipeline as the fixed experimental baseline.",
     "done", "Chunking → embeddings → ChromaDB, built and reused unchanged by every strategy."),
    ("Implement four distinct prompting strategies on top of it.",
     "done", "All four LCEL chains coded, unit-tested, and independently runnable."),
    ("Evaluate all strategies with an identical dataset and metric suite.",
     "progress", "RAGAS + a custom hallucination judge are built and validated; the full-scale run across all four strategies is in progress."),
    ("Compare trade-offs in accuracy, hallucination rate, and response quality.",
     "planned", "Reserved for the final review, once the full-scale run and statistical comparison are complete."),
]
y = 2.05
for text, status, note in objs:
    row_h = 1.02
    add_rect(s, Inches(0.6), Inches(y), Inches(12.1), Inches(row_h), fill=RGBColor(0xFA, 0xFB, 0xFC), line_color=LINE)
    kind = "done" if status == "done" else ("done" if status == "progress" else "planned")
    label = {"done": "DONE", "progress": "IN PROGRESS", "planned": "PLANNED"}[status]
    pill_fill = {"done": DONE_GREEN_BG, "progress": PLAN_AMBER_BG, "planned": RGBColor(0xEC, 0xEC, 0xEC)}[status]
    pill_ink = {"done": DONE_GREEN, "progress": PLAN_AMBER, "planned": MUTED}[status]
    pw = Inches(0.24 + 0.095 * len(label))
    pill = add_rect(s, Inches(0.78), Inches(y + 0.16), pw, Inches(0.3), fill=pill_fill, radius=0.5)
    ptf = pill.text_frame; ptf.margin_left = Inches(0.08); ptf.margin_right = Inches(0.08)
    ptf.margin_top = 0; ptf.margin_bottom = 0; ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pp = ptf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    pr = pp.add_run(); pr.text = label; pr.font.size = Pt(10); pr.font.bold = True; pr.font.name = SANS; pr.font.color.rgb = pill_ink
    add_textbox(s, Inches(0.78) + pw + Inches(0.22), Inches(y + 0.08), Inches(9.9), Inches(0.35),
                text, size=14.5, bold=True, color=INK, font=SANS)
    add_textbox(s, Inches(0.78) + pw + Inches(0.22), Inches(y + 0.5), Inches(10.8), Inches(0.45),
                note, size=12, color=INK_SOFT, font=SANS)
    y += row_h + 0.16
add_footer(s, 3)


# ===========================================================================
# SLIDE 4 -- System Architecture
# ===========================================================================
s = add_slide()
add_header(s, "03 · Architecture", "Fixed backbone, one variable stage")

def box(slide, cx, cy, w, h, text, sub=None, fill=WHITE, line=INK, text_color=INK, bold=True, size=12):
    left = Inches(cx - w / 2); top = Inches(cy - h / 2)
    shp = add_rect(slide, left, top, Inches(w), Inches(h), fill=fill, line_color=line, line_width=Pt(1.25), radius=0.12)
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.06); tf.margin_right = Inches(0.06)
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.name = SANS; r.font.color.rgb = text_color
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(9.5); r2.font.name = SANS; r2.font.color.rgb = INK_SOFT
    return shp

def connector(slide, x1, y1, x2, y2, color=INK_SOFT, weight=Pt(1.25), arrowhead=True):
    conn = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = weight
    if arrowhead:
        ln = conn.line._get_or_add_ln()
        tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
        ln.append(tail)
    return conn


# Layout: 4 evenly-spaced columns across x:[0.55, 10.15] (width 9.6), shared by
# both the fixed row and the variable row so everything lines up vertically.
COL_X = [1.70, 4.13, 6.57, 8.98]
BAND_LEFT, BAND_W = 0.55, 9.6
LLM_CX = (COL_X[0] + COL_X[-1]) / 2  # 5.34 -- centered under the 4 strategy boxes

# --- source corpus, feeding DOWN into the first fixed-pipeline box ---------
box(s, COL_X[0], 1.95, 2.7, 0.55, "RAGTruth QA subset", "5,934 rows · Hugging Face",
    fill=RGBColor(0xF5, 0xF7, 0xFA), line=LINE, bold=False, size=10.5)
connector(s, COL_X[0], 2.225, COL_X[0], 2.64, color=INK_SOFT)

# --- fixed band --------------------------------------------------------------
fixed_band = add_rect(s, Inches(BAND_LEFT), Inches(2.35), Inches(BAND_W), Inches(1.15), fill=None, line_color=NAVY, line_width=Pt(1.25), radius=0.08)
add_textbox(s, Inches(0.7), Inches(2.4), Inches(6), Inches(0.25), "FIXED — identical for every strategy", size=9.5, bold=True, color=NAVY, font=SANS)

fixed_labels = [("Chunk", "500 / 50 overlap"), ("Embed", "text-embedding-3-small"),
                ("Chroma DB", "vector store"), ("Retriever", "top-k = 4")]
for x, (label, sub) in zip(COL_X, fixed_labels):
    box(s, x, 2.98, 1.7, 0.62, label, sub)
for i in range(3):
    connector(s, COL_X[i] + 0.85, 2.98, COL_X[i + 1] - 0.85, 2.98)

# --- fan-out: retriever -> collector line -> each strategy box -------------
FANOUT_Y = 3.62
connector(s, COL_X[-1], 3.29, COL_X[-1], FANOUT_Y, arrowhead=False)
connector(s, COL_X[0], FANOUT_Y, COL_X[-1], FANOUT_Y, arrowhead=False)
for x in COL_X:
    connector(s, x, FANOUT_Y, x, 4.11)

# --- variable band -----------------------------------------------------------
var_band = add_rect(s, Inches(BAND_LEFT), Inches(3.85), Inches(BAND_W), Inches(1.05), fill=None, line_color=PINK_INK, line_width=Pt(1.25), radius=0.08)
add_textbox(s, Inches(0.7), Inches(3.9), Inches(7), Inches(0.25), "THE ONE VARIABLE — prompt strategy only", size=9.5, bold=True, color=PINK_INK, font=SANS)
colors4 = [RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34), RGBColor(0x1B, 0xAF, 0x7A), RGBColor(0x4A, 0x3A, 0xA7)]
labels4 = ["Zero-shot", "Few-shot", "Chain-of-Thought", "Structured Output"]
for x, lab, col in zip(COL_X, labels4, colors4):
    box(s, x, 4.42, 1.7, 0.62, lab, fill=WHITE, line=col, text_color=col, bold=True, size=10.5)

# --- fan-in: each strategy box -> collector line -> shared LLM call --------
FANIN_Y = 5.0
for x in COL_X:
    connector(s, x, 4.73, x, FANIN_Y, arrowhead=False)
connector(s, COL_X[0], FANIN_Y, COL_X[-1], FANIN_Y, arrowhead=False)
connector(s, LLM_CX, FANIN_Y, LLM_CX, 5.125)

llm = box(s, LLM_CX, 5.4, 3.2, 0.55, "gpt-4o-mini · T = 0.0", "shared generation call", fill=NAVY, line=NAVY, text_color=WHITE, size=12)
connector(s, LLM_CX, 5.675, LLM_CX, 5.9)

eval_box = box(s, LLM_CX, 6.25, 3.6, 0.55, "RAGAS + LLM Hallucination Judge", "calibrated against RAGTruth labels", fill=RGBColor(0xF5, 0xF7, 0xFA), line=LINE, size=11)

right_note = add_rect(s, Inches(10.35), Inches(3.75), Inches(2.4), Inches(2.8), fill=RGBColor(0xF5, 0xF7, 0xFA), line_color=LINE, radius=0.08)
tf = right_note.text_frame; tf.margin_left = Inches(0.18); tf.margin_top = Inches(0.18); tf.margin_right = Inches(0.16); tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run(); r.text = "Why this design"; r.font.bold = True; r.font.size = Pt(12.5); r.font.name = SANS; r.font.color.rgb = NAVY
p = tf.add_paragraph(); p.space_before = Pt(10)
r = p.add_run(); r.text = "Retrieval is built once and reused, read-only, by all four strategies — so any measured difference can only come from the prompt."
r.font.size = Pt(11); r.font.name = SANS; r.font.color.rgb = INK_SOFT
add_footer(s, 4)


# ===========================================================================
# SLIDE 5 -- Methodology / Control Variables Table
# ===========================================================================
s = add_slide()
add_header(s, "04 · Methodology", "What's fixed vs. what's tested")

rows = [
    ("Chunking", "RecursiveCharacterTextSplitter · 500 chars / 50 overlap", "Fixed"),
    ("Embeddings", "OpenAI text-embedding-3-small", "Fixed"),
    ("Vector store", "ChromaDB, persisted collection, built once", "Fixed"),
    ("Retrieval", "top-k = 4, same retriever object reused", "Fixed"),
    ("Generation model", "gpt-4o-mini, temperature 0.0", "Fixed"),
    ("Grading judge", "Same LLM judge & taxonomy for every strategy", "Fixed"),
    ("Prompt strategy", "Zero-shot / Few-shot / Chain-of-Thought / Structured Output", "Variable"),
]
tbl_left, tbl_top, tbl_w, tbl_h = Inches(0.6), Inches(2.1), Inches(12.1), Inches(4.5)
gtbl = s.shapes.add_table(len(rows) + 1, 3, tbl_left, tbl_top, tbl_w, tbl_h).table
gtbl.columns[0].width = Inches(2.6)
gtbl.columns[1].width = Inches(7.6)
gtbl.columns[2].width = Inches(1.9)
hdr = ["Stage", "Setting", "Status"]
for c, htext in enumerate(hdr):
    cell = gtbl.cell(0, c)
    cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
    cell.text_frame.paragraphs[0].text = htext
    r = cell.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(12.5); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = SANS
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
for ri, (stage, setting, status) in enumerate(rows, start=1):
    c0 = gtbl.cell(ri, 0); c0.text_frame.paragraphs[0].text = stage
    r0 = c0.text_frame.paragraphs[0].runs[0]; r0.font.size = Pt(12); r0.font.bold = True; r0.font.name = SANS; r0.font.color.rgb = INK
    c1 = gtbl.cell(ri, 1); c1.text_frame.paragraphs[0].text = setting
    r1 = c1.text_frame.paragraphs[0].runs[0]; r1.font.size = Pt(12); r1.font.name = SANS; r1.font.color.rgb = INK_SOFT
    c2 = gtbl.cell(ri, 2)
    c2.text_frame.paragraphs[0].text = status
    r2 = c2.text_frame.paragraphs[0].runs[0]; r2.font.size = Pt(11.5); r2.font.bold = True; r2.font.name = SANS
    r2.font.color.rgb = NAVY if status == "Fixed" else PINK_INK
    for c in range(3):
        cell = gtbl.cell(ri, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if ri % 2 else RGBColor(0xF7, 0xF9, 0xFA)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.12); cell.margin_top = Inches(0.05); cell.margin_bottom = Inches(0.05)
add_footer(s, 5)


# ===========================================================================
# SLIDE 6 -- Tech Stack
# ===========================================================================
s = add_slide()
add_header(s, "05 · Tech Stack", "What it's built with")

stack = [
    ("Orchestration", "LangChain (LCEL)"),
    ("Vector store", "ChromaDB"),
    ("Embeddings", "text-embedding-3-small"),
    ("Base LLM", "gpt-4o-mini"),
    ("Validation", "Pydantic v2"),
    ("Evaluation", "RAGAS"),
    ("Dataset", "RAGTruth (Hugging Face)"),
    ("Backend / GUI", "FastAPI + WebSocket + JS"),
    ("Testing", "pytest (14 tests)"),
]
cols3, rows3 = 3, 3
cw, ch = 3.85, 1.3
gx, gy = 0.35, 0.28
x0, y0 = 0.6, 2.15
for i, (role, name) in enumerate(stack):
    r_, c_ = divmod(i, cols3)
    x = x0 + c_ * (cw + gx)
    y = y0 + r_ * (ch + gy)
    card = add_rect(s, Inches(x), Inches(y), Inches(cw), Inches(ch), fill=RGBColor(0xF7, 0xF9, 0xFA), line_color=LINE, radius=0.1)
    tf = card.text_frame; tf.margin_left = Inches(0.22); tf.margin_top = Inches(0.2); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = role.upper(); r.font.size = Pt(10); r.font.bold = True; r.font.name = SANS; r.font.color.rgb = MUTED
    p2 = tf.add_paragraph(); p2.space_before = Pt(4)
    r2 = p2.add_run(); r2.text = name; r2.font.size = Pt(15.5); r2.font.bold = True; r2.font.name = SANS; r2.font.color.rgb = NAVY
add_footer(s, 6)


# ===========================================================================
# SLIDE 7 -- Implementation Status (the "50%" slide)
# ===========================================================================
s = add_slide()
add_header(s, "06 · Progress Checkpoint", "What's built vs. what's reserved for the final review")

done_items = [
    "RAGTruth data pipeline — loading, QA-subset filtering, eval/calibration split, corpus construction",
    "Chunking → embedding → ChromaDB indexing, shared read-only across all strategies",
    "All 4 prompting strategy chains implemented in LangChain LCEL",
    "Structured Output strategy's Pydantic schema (verbatim quotes + groundedness flag)",
    "RAGAS integration (Faithfulness, Answer Relevancy)",
    "Custom LLM hallucination judge, calibrated against RAGTruth's human labels",
    "Interactive web dashboard — live pipeline run, progress, and charts",
    "Automated test suite — 14 tests, no API key required",
]
planned_items = [
    "Full-scale benchmark run across all 4 strategies on the complete eval sample",
    "Statistical comparison of accuracy / hallucination rate / faithfulness across strategies",
    "Cost & latency trade-off analysis at scale",
    "Final recommendations: which strategy for which production scenario",
]

left_panel = add_rect(s, Inches(0.6), Inches(2.05), Inches(6.35), Inches(4.65), fill=DONE_GREEN_BG, line_color=None, radius=0.05)
add_textbox(s, Inches(0.85), Inches(2.22), Inches(5.8), Inches(0.35), "COMPLETED", size=13, bold=True, color=DONE_GREEN, font=SANS)
add_bullets(s, Inches(0.85), Inches(2.62), Inches(5.9), Inches(3.9), done_items, size=12.5,
            color=RGBColor(0x14, 0x40, 0x2A), marker="✓", marker_color=DONE_GREEN, gap_pt=9, line_spacing=1.08)

right_panel = add_rect(s, Inches(7.1), Inches(2.05), Inches(5.6), Inches(4.65), fill=PLAN_AMBER_BG, line_color=None, radius=0.05)
add_textbox(s, Inches(7.35), Inches(2.22), Inches(5.1), Inches(0.35), "RESERVED FOR FINAL REVIEW", size=13, bold=True, color=PLAN_AMBER, font=SANS)
add_bullets(s, Inches(7.35), Inches(2.62), Inches(5.1), Inches(3.9), planned_items, size=12.5,
            color=RGBColor(0x5A, 0x40, 0x08), marker="○", marker_color=PLAN_AMBER, gap_pt=11, line_spacing=1.15)
add_footer(s, 7)


# ===========================================================================
# SLIDE 8 -- Evaluation Framework Validation (safe to reveal)
# ===========================================================================
s = add_slide()
add_header(s, "07 · Framework Validation", "Is the judge itself trustworthy?",
           "Before grading any strategy, the LLM hallucination judge is checked against RAGTruth's own human annotations.")

why_box = add_rect(s, Inches(0.6), Inches(2.15), Inches(6.6), Inches(2.15), fill=RGBColor(0xF5, 0xF7, 0xFA), line_color=LINE, radius=0.06)
tf = why_box.text_frame; tf.margin_left = Inches(0.28); tf.margin_top = Inches(0.22); tf.margin_right = Inches(0.24); tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run(); r.text = "Why calibrate a judge?"; r.font.bold = True; r.font.size = Pt(14); r.font.name = SANS; r.font.color.rgb = NAVY
p = tf.add_paragraph(); p.space_before = Pt(10)
r = p.add_run()
r.text = ("RAGTruth's human labels were written for its own models' answers, not ours. "
          "So the judge is run on RAGTruth's original (context, response, human-label) "
          "triples, and its predictions are compared against those human labels — before it "
          "is trusted to grade any of our own strategies' output.")
r.font.size = Pt(12.5); r.font.name = SANS; r.font.color.rgb = INK_SOFT

metrics = [("Precision", "0.87"), ("Recall", "0.87"), ("F1", "0.87"), ("Accuracy", "0.87")]
mx0, my0, mw, mh, mg = 0.6, 4.55, 1.5, 1.15, 0.28
for i, (label, val) in enumerate(metrics):
    x = mx0 + i * (mw + mg)
    card = add_rect(s, Inches(x), Inches(my0), Inches(mw), Inches(mh), fill=NAVY, radius=0.12)
    tf = card.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = val; r.font.size = Pt(24); r.font.bold = True; r.font.name = SANS; r.font.color.rgb = WHITE
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = label.upper(); r2.font.size = Pt(9.5); r2.font.name = SANS; r2.font.color.rgb = RGBColor(0xC9, 0xD8, 0xE8)

note = add_rect(s, Inches(7.5), Inches(2.15), Inches(5.2), Inches(2.15), fill=WHITE, line_color=LINE, radius=0.06)
tf = note.text_frame; tf.margin_left = Inches(0.26); tf.margin_top = Inches(0.22); tf.margin_right = Inches(0.22); tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run(); r.text = "Methodology detail"; r.font.bold = True; r.font.size = Pt(13.5); r.font.name = SANS; r.font.color.rgb = NAVY
p = tf.add_paragraph(); p.space_before = Pt(8)
r = p.add_run(); r.text = "The judge is prompted to decompose each answer into atomic claims and check each one individually against the context — a single unsupported clause inside an otherwise-correct long answer must still be flagged."
r.font.size = Pt(11.5); r.font.name = SANS; r.font.color.rgb = INK_SOFT
add_textbox(s, Inches(0.6), Inches(6.0), Inches(11.5), Inches(0.4),
            "n = 30 calibration rows, held out from the evaluation set. These are framework-validation numbers — not the strategy comparison.",
            size=11, italic=True, color=MUTED, font=SANS)
add_footer(s, 8)


# ===========================================================================
# SLIDE 9 -- Live Demo Plan
# ===========================================================================
s = add_slide()
add_header(s, "08 · Live Demo", "What we'll show today")

demo_items = [
    ("Launch the web dashboard ", "— configure a small sample and click Run."),
    ("Watch retrieval happen live ", "— the same 4 chunks retrieved for every strategy, shown per question."),
    ("Watch one strategy answer, live ", "— the pipeline stepper, a scrolling feed of questions as they're answered and judged."),
    ("Show the calibration panel ", "— the judge-vs-human-label agreement numbers from Slide 8, computed live."),
    ("Browse an individual answer ", "— full retrieved context, generated answer, and judge rationale for one question."),
]
add_bullets(s, Inches(0.6), Inches(2.15), Inches(7.3), Inches(4.2), demo_items, size=15.5, gap_pt=16, line_spacing=1.15)

scope_box = add_rect(s, Inches(8.15), Inches(2.15), Inches(4.55), Inches(4.3), fill=PLAN_AMBER_BG, radius=0.08)
tf = scope_box.text_frame; tf.margin_left = Inches(0.28); tf.margin_top = Inches(0.24); tf.margin_right = Inches(0.24); tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run(); r.text = "Deliberately out of scope today"; r.font.bold = True; r.font.size = Pt(13.5); r.font.name = SANS; r.font.color.rgb = PLAN_AMBER
for item in [
    "The full-scale, all-4-strategies benchmark run",
    "Final accuracy / hallucination-rate comparison numbers",
    "Any conclusion about which strategy “wins”",
]:
    p = tf.add_paragraph(); p.space_before = Pt(12)
    r = p.add_run(); r.text = f"✗  {item}"; r.font.size = Pt(12.5); r.font.name = SANS; r.font.color.rgb = RGBColor(0x5A, 0x40, 0x08)
p = tf.add_paragraph(); p.space_before = Pt(18)
r = p.add_run(); r.text = "These are the headline result of the final review — showing them now would pre-empt it."
r.font.size = Pt(11.5); r.font.italic = True; r.font.name = SANS; r.font.color.rgb = RGBColor(0x5A, 0x40, 0x08)
add_footer(s, 9)


# ===========================================================================
# SLIDE 10 -- Next Steps / Plan for Final Review
# ===========================================================================
s = add_slide()
add_header(s, "09 · Next Steps", "Plan for the final review")

steps = [
    ("1", "Scale up the eval sample", "Run all 4 strategies across the full evaluation set (n ≥ 60 questions) for statistically meaningful gaps."),
    ("2", "Full comparative analysis", "Accuracy, hallucination rate, RAGAS Faithfulness / Answer Relevancy, latency, and cost — compared strategy vs. strategy."),
    ("3", "Statistical robustness check", "Quantify how much the sample size affects confidence in the observed differences."),
    ("4", "Final report & recommendation", "Which strategy to use for which production RAG scenario, with the trade-offs made explicit."),
]
y = 2.15
for num, title, desc in steps:
    row_h = 1.02
    circle = add_rect(s, Inches(0.6), Inches(y + 0.08), Inches(0.55), Inches(0.55), fill=NAVY, radius=0.5)
    ctf = circle.text_frame; ctf.vertical_anchor = MSO_ANCHOR.MIDDLE; ctf.margin_left = 0; ctf.margin_right = 0
    cp = ctf.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
    cr = cp.add_run(); cr.text = num; cr.font.size = Pt(16); cr.font.bold = True; cr.font.name = SANS; cr.font.color.rgb = WHITE
    add_textbox(s, Inches(1.4), Inches(y), Inches(10.8), Inches(0.35), title, size=15.5, bold=True, color=INK, font=SANS)
    add_textbox(s, Inches(1.4), Inches(y + 0.38), Inches(10.8), Inches(0.5), desc, size=12.5, color=INK_SOFT, font=SANS, line_spacing=1.15)
    y += row_h + 0.28
add_footer(s, 10)


prs.save(str(OUT_PATH))
print(f"Saved -> {OUT_PATH}")

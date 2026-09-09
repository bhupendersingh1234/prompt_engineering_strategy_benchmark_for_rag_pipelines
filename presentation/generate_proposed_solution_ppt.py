"""
Generates the formal "Proposed Solution & 50% Work Completion" review deck
(11 slides, 10-minute presentation) for "Prompt Engineering Strategy
Benchmark for RAG Pipelines", matching the panel's required structure:
signed/scanned title, problem + background, literature survey (10 refs),
existing solutions + limitations, proposed solution/scope/methodology,
work completed (~50%), project timeline, tools & technologies.

All 10 literature references were verified against real sources (arXiv /
ACL Anthology / ACM DL) via web search during authoring -- not recalled
from memory alone. See the LITERATURE list below for each paper's venue.

Like generate_first_review_ppt.py, this deck deliberately stops short of
the full 4-strategy comparison numbers -- those are reserved for the
final review. The architecture diagram, control-variables table, tech
stack, and judge-calibration slides are carried over unchanged from that
already-verified deck (same geometry, re-checked here for this script's
own header/footer offsets).

Usage:
    python presentation/generate_proposed_solution_ppt.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

OUT_PATH = Path(__file__).resolve().parent / "Proposed_Solution_and_50_Percent_Completion.pptx"

# ---------------------------------------------------------------------------
# Palette / type system (identical to generate_first_review_ppt.py, so the
# two decks read as one visual identity)
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x0D, 0x3B, 0x66)
INK = RGBColor(0x1A, 0x1D, 0x22)
INK_SOFT = RGBColor(0x4A, 0x4F, 0x58)
MUTED = RGBColor(0x8A, 0x8F, 0x98)
LINE = RGBColor(0xD8, 0xDC, 0xE2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAGE_BG = RGBColor(0xFF, 0xFF, 0xFF)
PINK_BG = RGBColor(0xFB, 0xE4, 0xE4)
PINK_INK = RGBColor(0xA9, 0x33, 0x3F)
DONE_GREEN = RGBColor(0x1E, 0x7A, 0x4C)
DONE_GREEN_BG = RGBColor(0xE3, 0xF3, 0xE9)
PLAN_AMBER = RGBColor(0x9A, 0x6A, 0x12)
PLAN_AMBER_BG = RGBColor(0xFB, 0xF0, 0xDE)
GREY_BG = RGBColor(0xF0, 0xF1, 0xF3)
GREY_INK = RGBColor(0x5A, 0x5E, 0x66)

SERIF = "Georgia"
SANS = "Segoe UI"
MONO = "Consolas"

REVIEW_LABEL = "First Review"
FOOTER_TEXT = f"B.Tech Project – I | {REVIEW_LABEL} | School of Computer Science and Engineering | VIT Chennai"

TEAM = [
    ("23BAI1057", "Astitva Thakkar"),
    ("23BAI1528", "Parth Sharma"),
    ("23BRS1107", "Bhupender Singh"),
]
FACULTY = {
    "Name": "Manmohan Sharma",
    "ERP ID": "53622",
    "School": "SENSE (School of Electronics and Communication Engineering)",
}

# ---------------------------------------------------------------------------
# Literature survey -- 10 references, each verified via web search against
# arXiv / ACL Anthology / ACM Digital Library during authoring (not
# recalled from memory alone, given citation accuracy matters for an
# academic submission).
# ---------------------------------------------------------------------------
LITERATURE = [
    ("Lewis et al. (2020, NeurIPS)", "Introduced Retrieval-Augmented Generation (RAG)"),
    ("Brown et al. (2020, NeurIPS)", "GPT-3; established few-shot in-context learning"),
    ("Shuster et al. (2021, EMNLP Findings)", "Retrieval augmentation reduces hallucination in dialogue"),
    ("Wei et al. (2022, NeurIPS)", "Chain-of-Thought prompting improves multi-step reasoning"),
    ("Kojima et al. (2022, NeurIPS)", "Zero-shot CoT via “Let’s think step by step”"),
    ("Liu et al. (2023, ACM Comput. Surv.)", "Systematic survey of prompting methods in NLP"),
    ("Ji et al. (2023, ACM Comput. Surv.)", "Survey of hallucination in natural language generation"),
    ("Gao et al. (2023, arXiv survey)", "Survey of RAG techniques for large language models"),
    ("Es et al. (2024, EACL)", "RAGAS: reference-free evaluation metrics for RAG"),
    ("Niu et al. (2024, ACL)", "RAGTruth: hallucination corpus used as this project’s dataset"),
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ===========================================================================
# Shared helpers (same as generate_first_review_ppt.py)
# ===========================================================================
def add_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAGE_BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    spTree = slide.shapes._spTree
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)
    return slide


def add_textbox(slide, left, top, width, height, text, size=14, color=INK,
                 bold=False, italic=False, font=SANS, align=PP_ALIGN.LEFT,
                 anchor=MSO_ANCHOR.TOP, line_spacing=1.0, wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return box


def add_bullets(slide, left, top, width, height, items, size=15, color=INK,
                 font=SANS, gap_pt=10, marker="—", marker_color=None,
                 line_spacing=1.12):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(gap_pt)
        r0 = p.add_run()
        r0.text = f"{marker}  "
        r0.font.size = Pt(size)
        r0.font.name = font
        r0.font.color.rgb = marker_color or color
        r0.font.bold = True
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run()
            r1.text = lead
            r1.font.size = Pt(size)
            r1.font.name = font
            r1.font.bold = True
            r1.font.color.rgb = color
            if rest:
                r2 = p.add_run()
                r2.text = rest
                r2.font.size = Pt(size)
                r2.font.name = font
                r2.font.color.rgb = color
        else:
            r1 = p.add_run()
            r1.text = item
            r1.font.size = Pt(size)
            r1.font.name = font
            r1.font.color.rgb = color
    return box


def add_rect(slide, left, top, width, height, fill=None, line_color=None,
             line_width=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    if radius:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is not None:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    if line_color is not None:
        shp.line.color.rgb = line_color
        shp.line.width = line_width or Pt(0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_dashed_rect(slide, left, top, width, height, line_color=MUTED):
    shp = add_rect(slide, left, top, width, height, fill=None, line_color=line_color, line_width=Pt(1))
    ln = shp.line._get_or_add_ln()
    dash = ln.makeelement(qn('a:prstDash'), {'val': 'dash'})
    ln.append(dash)
    return shp


def add_line(slide, left, top, width, color=LINE, weight=Pt(1)):
    ln = slide.shapes.add_connector(1, left, top, left + width, top)
    ln.line.color.rgb = color
    ln.line.width = weight
    return ln


def add_footer(slide, page_no):
    bar = add_rect(slide, Inches(0.5), Inches(6.95), Inches(12.33), Inches(0.4), fill=PINK_BG)
    tf = bar.text_frame
    tf.margin_left = Inches(0.18)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = f"⚑  {FOOTER_TEXT}"
    r.font.size = Pt(10.5)
    r.font.name = SERIF
    r.font.color.rgb = PINK_INK
    add_textbox(slide, Inches(12.0), Inches(6.98), Inches(0.7), Inches(0.32),
                f"{page_no:02d}", size=10.5, color=PINK_INK, font=MONO,
                align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    return bar


def add_header(slide, eyebrow, title, subtitle=None):
    add_textbox(slide, Inches(0.6), Inches(0.35), Inches(11.5), Inches(0.32),
                eyebrow.upper(), size=11.5, color=NAVY, bold=True, font=SANS)
    add_textbox(slide, Inches(0.6), Inches(0.68), Inches(11.8), Inches(0.75),
                title, size=27, color=INK, bold=False, font=SERIF)
    y = Inches(1.42)
    if subtitle:
        add_textbox(slide, Inches(0.6), y, Inches(11.8), Inches(0.5),
                    subtitle, size=13, color=INK_SOFT, font=SANS, line_spacing=1.2)
        y = Inches(1.85)
    add_line(slide, Inches(0.6), Inches(1.34) if not subtitle else Inches(1.78),
             Inches(12.1), color=LINE, weight=Pt(1))
    return y


def box(slide, cx, cy, w, h, text, sub=None, fill=WHITE, line=INK, text_color=INK, bold=True, size=12):
    left = Inches(cx - w / 2); top = Inches(cy - h / 2)
    shp = add_rect(slide, left, top, Inches(w), Inches(h), fill=fill, line_color=line, line_width=Pt(1.25), radius=0.12)
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.06); tf.margin_right = Inches(0.06)
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.name = SANS; r.font.color.rgb = text_color
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(9.5); r2.font.name = SANS; r2.font.color.rgb = INK_SOFT
    return shp


def connector(slide, x1, y1, x2, y2, color=INK_SOFT, weight=Pt(1.25), arrowhead=True):
    conn = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = weight
    if arrowhead:
        ln = conn.line._get_or_add_ln()
        tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
        ln.append(tail)
    return conn


def make_table(slide, left, top, w, h, col_widths, header, rows_data, font_size=11.5,
               header_size=12, row_colors=True):
    n_rows = len(rows_data) + 1
    n_cols = len(header)
    tbl = slide.shapes.add_table(n_rows, n_cols, left, top, w, h).table
    for i, cw in enumerate(col_widths):
        tbl.columns[i].width = cw
    for c, htext in enumerate(header):
        cell = tbl.cell(0, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        cell.text_frame.paragraphs[0].text = htext
        r = cell.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(header_size); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = SANS
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.12)
    for ri, row_vals in enumerate(rows_data, start=1):
        for ci, val in enumerate(row_vals):
            cell = tbl.cell(ri, ci)
            cell.text_frame.paragraphs[0].text = str(val)
            r = cell.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(font_size); r.font.name = SANS
            r.font.bold = (ci == 0)
            r.font.color.rgb = INK if ci == 0 else INK_SOFT
            cell.fill.solid()
            cell.fill.fore_color.rgb = (WHITE if not row_colors else (WHITE if ri % 2 else RGBColor(0xF7, 0xF9, 0xFA)))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.12); cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
    return tbl


# ===========================================================================
# SLIDE 1 -- Title (signed & scanned)
# ===========================================================================
s = add_slide()
add_textbox(s, Inches(0), Inches(0.68), SW, Inches(0.4),
            f"VIT CHENNAI  —  B.Tech Project – I  —  {REVIEW_LABEL}",
            size=14, color=INK_SOFT, font=SERIF, align=PP_ALIGN.CENTER)
add_textbox(s, Inches(0.8), Inches(1.05), Inches(11.73), Inches(0.55),
            "Proposed Solution & 50% Work Completion", size=15, italic=True,
            color=NAVY, font=SANS, align=PP_ALIGN.CENTER)
add_textbox(s, Inches(0.8), Inches(1.48), Inches(11.73), Inches(1.0),
            "Prompt Engineering Strategy Benchmark for RAG Pipelines",
            size=32, color=INK, font=SERIF, align=PP_ALIGN.CENTER)
add_line(s, Inches(0.45), Inches(2.55), Inches(12.43), color=LINE, weight=Pt(1))

team_box = add_rect(s, Inches(0.45), Inches(2.85), Inches(5.75), Inches(2.0), fill=NAVY)
tf = team_box.text_frame
tf.margin_left = Inches(0.28); tf.margin_top = Inches(0.2); tf.margin_right = Inches(0.2)
p = tf.paragraphs[0]
r = p.add_run(); r.text = "Project Team"; r.font.size = Pt(15); r.font.name = SERIF; r.font.color.rgb = WHITE
for reg, name in TEAM:
    p = tf.add_paragraph(); p.space_before = Pt(12)
    r1 = p.add_run(); r1.text = reg; r1.font.bold = True; r1.font.size = Pt(13); r1.font.name = SANS; r1.font.color.rgb = WHITE
    r2 = p.add_run(); r2.text = f"  —  {name}"; r2.font.size = Pt(13); r2.font.name = SANS; r2.font.color.rgb = RGBColor(0xDD, 0xE6, 0xF0)

fac_left = Inches(6.5)
add_textbox(s, fac_left, Inches(2.85), Inches(5.3), Inches(0.35), "Faculty Guide", size=15, color=INK, font=SERIF)
fy = 3.32
for k, v in FACULTY.items():
    add_textbox(s, fac_left, Inches(fy), Inches(5.3), Inches(0.32), f"{k}: {v}", size=12.5, color=INK_SOFT, font=SANS)
    fy += 0.37
add_textbox(s, fac_left, Inches(fy + 0.08), Inches(2.0), Inches(0.3), "Signature:", size=12.5, color=INK_SOFT, font=SANS)
add_dashed_rect(s, fac_left + Inches(1.35), Inches(fy - 0.02), Inches(3.9), Inches(0.55))
add_textbox(s, fac_left + Inches(1.5), Inches(fy + 0.13), Inches(3.6), Inches(0.3),
            "insert scanned signature image here", size=9.5, italic=True, color=MUTED, font=SANS)

footer1 = add_rect(s, Inches(0.45), Inches(6.55), Inches(12.43), Inches(0.55), fill=PINK_BG)
tf = footer1.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.22)
p = tf.paragraphs[0]
r = p.add_run(); r.text = f"⚑  {FOOTER_TEXT}"; r.font.size = Pt(12); r.font.name = SERIF; r.font.color.rgb = PINK_INK


# ===========================================================================
# SLIDE 2 -- Problem Statement & Background
# ===========================================================================
s = add_slide()
add_header(s, "01 · Problem & Background", "Retrieval fixes hallucination — prompting still decides how well")

add_bullets(s, Inches(0.6), Inches(2.1), Inches(6.9), Inches(4.3), [
    ("Background: ", "Large Language Models generate fluent but sometimes fabricated text when answering from parametric knowledge alone (Ji et al., 2023)."),
    ("Retrieval-Augmented Generation ", "(Lewis et al., 2020) mitigates this by grounding generation in retrieved documents, and is now the dominant approach for building trustworthy LLM applications."),
    ("The open problem: ", "answer quality inside a RAG pipeline still depends heavily on how the retrieved context is handed to the LLM — the prompting strategy — not just on retrieval quality itself."),
    ("The gap: ", "no controlled study isolates prompting strategy as the sole variable while holding chunking, embeddings, retrieval, and the base model fixed."),
], size=14.5, gap_pt=14)

card = add_rect(s, Inches(7.85), Inches(2.1), Inches(4.9), Inches(4.3), fill=RGBColor(0xF5, 0xF7, 0xFA), line_color=LINE)
tf = card.text_frame
tf.margin_left = Inches(0.32); tf.margin_top = Inches(0.28); tf.margin_right = Inches(0.28)
p = tf.paragraphs[0]
r = p.add_run(); r.text = "Problem statement"; r.font.bold = True; r.font.size = Pt(14); r.font.name = SANS; r.font.color.rgb = NAVY
p = tf.add_paragraph(); p.space_before = Pt(10)
r = p.add_run()
r.text = ("Systematically benchmark four prompting strategies — Zero-shot, Few-shot, "
          "Chain-of-Thought, and Structured Output — on an identical RAG pipeline, "
          "isolating prompt design as the only variable, to identify the most "
          "reliable, grounded approach for production RAG systems.")
r.font.size = Pt(13); r.font.name = SANS; r.font.color.rgb = INK
p = tf.add_paragraph(); p.space_before = Pt(16)
r = p.add_run(); r.text = "Domain: AI · LLMs · RAG · NLP"; r.font.size = Pt(11.5); r.font.italic = True; r.font.name = SANS; r.font.color.rgb = MUTED
add_footer(s, 2)


# ===========================================================================
# SLIDE 3 -- Literature Survey (10 references)
# ===========================================================================
s = add_slide()
add_header(s, "02 · Literature Survey", "Ten references spanning RAG, prompting, and hallucination evaluation")

lit_rows = [(str(i + 1), ref, contrib) for i, (ref, contrib) in enumerate(LITERATURE)]
make_table(
    s, Inches(0.6), Inches(2.05), Inches(12.1), Inches(4.35),
    [Inches(0.5), Inches(3.7), Inches(7.9)],
    ["#", "Reference", "Key Contribution"],
    lit_rows, font_size=11, header_size=12,
)
add_footer(s, 3)


# ===========================================================================
# SLIDE 4 -- Existing Solutions & Their Limitations
# ===========================================================================
s = add_slide()
add_header(s, "03 · Existing Solutions", "What's already out there, and where it falls short")

existing_rows = [
    ("Standard RAG pipelines", "Retrieval reduces hallucination, but prompt design is treated as an implementation detail, not studied systematically."),
    ("Few-shot / CoT prompting research", "Extensively studied for general reasoning tasks; rarely evaluated inside a RAG pipeline with retrieval held constant."),
    ("RAG evaluation frameworks (RAGAS)", "Provide faithfulness / relevancy metrics, but don't benchmark prompting strategies against each other."),
    ("Hallucination corpora (RAGTruth)", "Rich human-labeled hallucination data — but labels apply to the original models' answers, not a controlled ablation."),
    ("Production RAG systems (industry)", "Prompting strategy is usually chosen ad hoc, without a systematic accuracy-vs-hallucination-vs-cost comparison."),
]
make_table(
    s, Inches(0.6), Inches(2.05), Inches(12.1), Inches(3.15),
    [Inches(4.0), Inches(8.1)],
    ["Existing Approach", "Limitation"],
    existing_rows, font_size=12,
)

gap_box = add_rect(s, Inches(0.6), Inches(5.4), Inches(12.1), Inches(1.15), fill=PINK_BG, radius=0.06)
tf = gap_box.text_frame
tf.margin_left = Inches(0.3); tf.margin_top = Inches(0.16); tf.margin_right = Inches(0.3); tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run(); r.text = "The gap this project addresses:  "; r.font.bold = True; r.font.size = Pt(13); r.font.name = SANS; r.font.color.rgb = PINK_INK
r2 = p.add_run()
r2.text = ("No existing work isolates prompting strategy as the sole variable in a RAG "
           "pipeline while holding retrieval, embeddings, and the base model fixed.")
r2.font.size = Pt(13); r2.font.name = SANS; r2.font.color.rgb = RGBColor(0x6B, 0x24, 0x2E)
add_footer(s, 4)


# ===========================================================================
# SLIDE 5 -- Proposed Solution & Scope
# ===========================================================================
s = add_slide()
add_header(s, "04 · Proposed Solution", "One fixed pipeline, four prompting strategies, controlled comparison")

add_bullets(s, Inches(0.6), Inches(2.1), Inches(6.9), Inches(2.0), [
    ("Fix everything except the prompt: ", "chunking, embeddings, vector store, top-k retrieval, base LLM, and temperature are held identical across every strategy."),
    ("Vary only the prompting strategy: ", "Zero-shot, Few-shot, Chain-of-Thought, and Structured Output, each implemented as an independent LangChain LCEL chain."),
    ("Grade with a validated framework: ", "RAGAS (Faithfulness, Answer Relevancy) plus a custom LLM judge calibrated against RAGTruth's own human hallucination labels."),
], size=14, gap_pt=14)

scope_in = add_rect(s, Inches(7.85), Inches(2.1), Inches(4.9), Inches(2.05), fill=DONE_GREEN_BG, radius=0.06)
tf = scope_in.text_frame; tf.margin_left = Inches(0.26); tf.margin_top = Inches(0.18); tf.margin_right = Inches(0.2); tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run(); r.text = "In scope"; r.font.bold = True; r.font.size = Pt(13); r.font.name = SANS; r.font.color.rgb = DONE_GREEN
for item in ["RAGTruth QA subset (Hugging Face)", "4 prompting strategies", "ChromaDB + text-embedding-3-small", "gpt-4o-mini generation & judging", "RAGAS + custom hallucination judge"]:
    p = tf.add_paragraph(); p.space_before = Pt(6)
    r = p.add_run(); r.text = f"✓  {item}"; r.font.size = Pt(11.5); r.font.name = SANS; r.font.color.rgb = RGBColor(0x14, 0x40, 0x2A)

scope_out = add_rect(s, Inches(7.85), Inches(4.25), Inches(4.9), Inches(2.05), fill=GREY_BG, radius=0.06)
tf = scope_out.text_frame; tf.margin_left = Inches(0.26); tf.margin_top = Inches(0.18); tf.margin_right = Inches(0.2); tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run(); r.text = "Out of scope"; r.font.bold = True; r.font.size = Pt(13); r.font.name = SANS; r.font.color.rgb = GREY_INK
for item in ["Other embedding models / vector DBs", "Other LLM providers / open-weight models", "Fine-tuning or RLHF", "Datasets beyond RAGTruth", "Multi-turn conversation"]:
    p = tf.add_paragraph(); p.space_before = Pt(6)
    r = p.add_run(); r.text = f"✗  {item}"; r.font.size = Pt(11.5); r.font.name = SANS; r.font.color.rgb = GREY_INK

add_textbox(s, Inches(0.6), Inches(4.4), Inches(6.9), Inches(1.9),
            "Methodology in detail — architecture, control-variable table, and framework "
            "validation — is covered on the following slides.",
            size=12.5, italic=True, color=MUTED, font=SANS, line_spacing=1.3)
add_footer(s, 5)


# ===========================================================================
# SLIDE 6 -- Methodology: System Architecture (verified diagram, reused)
# ===========================================================================
s = add_slide()
add_header(s, "05 · Methodology", "System architecture — fixed backbone, one variable stage")

COL_X = [1.70, 4.13, 6.57, 8.98]
BAND_LEFT, BAND_W = 0.55, 9.6
LLM_CX = (COL_X[0] + COL_X[-1]) / 2

box(s, COL_X[0], 1.95, 2.7, 0.55, "RAGTruth QA subset", "5,934 rows · Hugging Face",
    fill=RGBColor(0xF5, 0xF7, 0xFA), line=LINE, bold=False, size=10.5)
connector(s, COL_X[0], 2.225, COL_X[0], 2.64, color=INK_SOFT)

fixed_band = add_rect(s, Inches(BAND_LEFT), Inches(2.35), Inches(BAND_W), Inches(1.15), fill=None, line_color=NAVY, line_width=Pt(1.25), radius=0.08)
add_textbox(s, Inches(0.7), Inches(2.4), Inches(6), Inches(0.25), "FIXED — identical for every strategy", size=9.5, bold=True, color=NAVY, font=SANS)

fixed_labels = [("Chunk", "500 / 50 overlap"), ("Embed", "text-embedding-3-small"),
                ("Chroma DB", "vector store"), ("Retriever", "top-k = 4")]
for x, (label, sub) in zip(COL_X, fixed_labels):
    box(s, x, 2.98, 1.7, 0.62, label, sub)
for i in range(3):
    connector(s, COL_X[i] + 0.85, 2.98, COL_X[i + 1] - 0.85, 2.98)

FANOUT_Y = 3.62
connector(s, COL_X[-1], 3.29, COL_X[-1], FANOUT_Y, arrowhead=False)
connector(s, COL_X[0], FANOUT_Y, COL_X[-1], FANOUT_Y, arrowhead=False)
for x in COL_X:
    connector(s, x, FANOUT_Y, x, 4.11)

var_band = add_rect(s, Inches(BAND_LEFT), Inches(3.85), Inches(BAND_W), Inches(1.05), fill=None, line_color=PINK_INK, line_width=Pt(1.25), radius=0.08)
add_textbox(s, Inches(0.7), Inches(3.9), Inches(7), Inches(0.25), "THE ONE VARIABLE — prompt strategy only", size=9.5, bold=True, color=PINK_INK, font=SANS)
colors4 = [RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34), RGBColor(0x1B, 0xAF, 0x7A), RGBColor(0x4A, 0x3A, 0xA7)]
labels4 = ["Zero-shot", "Few-shot", "Chain-of-Thought", "Structured Output"]
for x, lab, col in zip(COL_X, labels4, colors4):
    box(s, x, 4.42, 1.7, 0.62, lab, fill=WHITE, line=col, text_color=col, bold=True, size=10.5)

FANIN_Y = 5.0
for x in COL_X:
    connector(s, x, 4.73, x, FANIN_Y, arrowhead=False)
connector(s, COL_X[0], FANIN_Y, COL_X[-1], FANIN_Y, arrowhead=False)
connector(s, LLM_CX, FANIN_Y, LLM_CX, 5.125)

box(s, LLM_CX, 5.4, 3.2, 0.55, "gpt-4o-mini · T = 0.0", "shared generation call", fill=NAVY, line=NAVY, text_color=WHITE, size=12)
connector(s, LLM_CX, 5.675, LLM_CX, 5.9)

box(s, LLM_CX, 6.25, 3.6, 0.55, "RAGAS + LLM Hallucination Judge", "calibrated against RAGTruth labels", fill=RGBColor(0xF5, 0xF7, 0xFA), line=LINE, size=11)

right_note = add_rect(s, Inches(10.35), Inches(3.75), Inches(2.4), Inches(2.8), fill=RGBColor(0xF5, 0xF7, 0xFA), line_color=LINE, radius=0.08)
tf = right_note.text_frame; tf.margin_left = Inches(0.18); tf.margin_top = Inches(0.18); tf.margin_right = Inches(0.16); tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run(); r.text = "Why this design"; r.font.bold = True; r.font.size = Pt(12.5); r.font.name = SANS; r.font.color.rgb = NAVY
p = tf.add_paragraph(); p.space_before = Pt(10)
r = p.add_run(); r.text = "Retrieval is built once and reused, read-only, by all four strategies — so any measured difference can only come from the prompt."
r.font.size = Pt(11); r.font.name = SANS; r.font.color.rgb = INK_SOFT
add_footer(s, 6)


# ===========================================================================
# SLIDE 7 -- Methodology: Control Variables Table (reused)
# ===========================================================================
s = add_slide()
add_header(s, "05 · Methodology", "What's fixed vs. what's tested")

cv_rows = [
    ("Chunking", "RecursiveCharacterTextSplitter · 500 chars / 50 overlap", "Fixed"),
    ("Embeddings", "OpenAI text-embedding-3-small", "Fixed"),
    ("Vector store", "ChromaDB, persisted collection, built once", "Fixed"),
    ("Retrieval", "top-k = 4, same retriever object reused", "Fixed"),
    ("Generation model", "gpt-4o-mini, temperature 0.0", "Fixed"),
    ("Grading judge", "Same LLM judge & taxonomy for every strategy", "Fixed"),
    ("Prompt strategy", "Zero-shot / Few-shot / Chain-of-Thought / Structured Output", "Variable"),
]
tbl_left, tbl_top, tbl_w, tbl_h = Inches(0.6), Inches(2.1), Inches(12.1), Inches(4.5)
gtbl = s.shapes.add_table(len(cv_rows) + 1, 3, tbl_left, tbl_top, tbl_w, tbl_h).table
gtbl.columns[0].width = Inches(2.6)
gtbl.columns[1].width = Inches(7.6)
gtbl.columns[2].width = Inches(1.9)
for c, htext in enumerate(["Stage", "Setting", "Status"]):
    cell = gtbl.cell(0, c)
    cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
    cell.text_frame.paragraphs[0].text = htext
    r = cell.text_frame.paragraphs[0].runs[0]
    r.font.size = Pt(12.5); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = SANS
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
for ri, (stage, setting, status) in enumerate(cv_rows, start=1):
    c0 = gtbl.cell(ri, 0); c0.text_frame.paragraphs[0].text = stage
    r0 = c0.text_frame.paragraphs[0].runs[0]; r0.font.size = Pt(12); r0.font.bold = True; r0.font.name = SANS; r0.font.color.rgb = INK
    c1 = gtbl.cell(ri, 1); c1.text_frame.paragraphs[0].text = setting
    r1 = c1.text_frame.paragraphs[0].runs[0]; r1.font.size = Pt(12); r1.font.name = SANS; r1.font.color.rgb = INK_SOFT
    c2 = gtbl.cell(ri, 2); c2.text_frame.paragraphs[0].text = status
    r2 = c2.text_frame.paragraphs[0].runs[0]; r2.font.size = Pt(11.5); r2.font.bold = True; r2.font.name = SANS
    r2.font.color.rgb = NAVY if status == "Fixed" else PINK_INK
    for c in range(3):
        cell = gtbl.cell(ri, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if ri % 2 else RGBColor(0xF7, 0xF9, 0xFA)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.12); cell.margin_top = Inches(0.05); cell.margin_bottom = Inches(0.05)
add_footer(s, 7)


# ===========================================================================
# SLIDE 8 -- Work Completed (~50%)
# ===========================================================================
s = add_slide()
add_header(s, "06 · Work Completed", "Approximately 50% — infrastructure and framework validated")

done_items = [
    "RAGTruth data pipeline — loading, QA-subset filtering, eval/calibration split, corpus construction",
    "Chunking → embedding → ChromaDB indexing, shared read-only across all strategies",
    "All 4 prompting strategy chains implemented in LangChain LCEL",
    "Structured Output strategy's Pydantic schema (verbatim quotes + groundedness flag)",
    "RAGAS integration (Faithfulness, Answer Relevancy)",
    "Custom LLM hallucination judge, calibrated against RAGTruth's human labels",
    "Interactive web dashboard — live pipeline run, progress, and charts",
    "Automated test suite — 14 tests, no API key required",
]
planned_items = [
    "Full-scale benchmark run across all 4 strategies on the complete eval sample",
    "Statistical comparison of accuracy / hallucination rate / faithfulness across strategies",
    "Cost & latency trade-off analysis at scale",
    "Final recommendations: which strategy for which production scenario",
]

left_panel = add_rect(s, Inches(0.6), Inches(2.05), Inches(6.35), Inches(4.65), fill=DONE_GREEN_BG, line_color=None, radius=0.05)
add_textbox(s, Inches(0.85), Inches(2.22), Inches(5.8), Inches(0.35), "COMPLETED", size=13, bold=True, color=DONE_GREEN, font=SANS)
add_bullets(s, Inches(0.85), Inches(2.62), Inches(5.9), Inches(3.9), done_items, size=12.5,
            color=RGBColor(0x14, 0x40, 0x2A), marker="✓", marker_color=DONE_GREEN, gap_pt=9, line_spacing=1.08)

right_panel = add_rect(s, Inches(7.1), Inches(2.05), Inches(5.6), Inches(4.65), fill=PLAN_AMBER_BG, line_color=None, radius=0.05)
add_textbox(s, Inches(7.35), Inches(2.22), Inches(5.1), Inches(0.35), "REMAINING FOR NEXT REVIEWS", size=13, bold=True, color=PLAN_AMBER, font=SANS)
add_bullets(s, Inches(7.35), Inches(2.62), Inches(5.1), Inches(3.9), planned_items, size=12.5,
            color=RGBColor(0x5A, 0x40, 0x08), marker="○", marker_color=PLAN_AMBER, gap_pt=11, line_spacing=1.15)
add_footer(s, 8)


# ===========================================================================
# SLIDE 9 -- Evaluation Framework Validation (reused)
# ===========================================================================
s = add_slide()
add_header(s, "06 · Work Completed", "Is the judge itself trustworthy?",
           "Before grading any strategy, the LLM hallucination judge is checked against RAGTruth's own human annotations.")

why_box = add_rect(s, Inches(0.6), Inches(2.15), Inches(6.6), Inches(2.15), fill=RGBColor(0xF5, 0xF7, 0xFA), line_color=LINE, radius=0.06)
tf = why_box.text_frame; tf.margin_left = Inches(0.28); tf.margin_top = Inches(0.22); tf.margin_right = Inches(0.24); tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run(); r.text = "Why calibrate a judge?"; r.font.bold = True; r.font.size = Pt(14); r.font.name = SANS; r.font.color.rgb = NAVY
p = tf.add_paragraph(); p.space_before = Pt(10)
r = p.add_run()
r.text = ("RAGTruth's human labels were written for its own models' answers, not ours. "
          "So the judge is run on RAGTruth's original (context, response, human-label) "
          "triples, and its predictions are compared against those human labels — before it "
          "is trusted to grade any of our own strategies' output.")
r.font.size = Pt(12.5); r.font.name = SANS; r.font.color.rgb = INK_SOFT

metrics = [("Precision", "0.87"), ("Recall", "0.87"), ("F1", "0.87"), ("Accuracy", "0.87")]
mx0, my0, mw, mh, mg = 0.6, 4.55, 1.5, 1.15, 0.28
for i, (label, val) in enumerate(metrics):
    x = mx0 + i * (mw + mg)
    card = add_rect(s, Inches(x), Inches(my0), Inches(mw), Inches(mh), fill=NAVY, radius=0.12)
    tf = card.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = val; r.font.size = Pt(24); r.font.bold = True; r.font.name = SANS; r.font.color.rgb = WHITE
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = label.upper(); r2.font.size = Pt(9.5); r2.font.name = SANS; r2.font.color.rgb = RGBColor(0xC9, 0xD8, 0xE8)

note = add_rect(s, Inches(7.5), Inches(2.15), Inches(5.2), Inches(2.15), fill=WHITE, line_color=LINE, radius=0.06)
tf = note.text_frame; tf.margin_left = Inches(0.26); tf.margin_top = Inches(0.22); tf.margin_right = Inches(0.22); tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run(); r.text = "Methodology detail"; r.font.bold = True; r.font.size = Pt(13.5); r.font.name = SANS; r.font.color.rgb = NAVY
p = tf.add_paragraph(); p.space_before = Pt(8)
r = p.add_run(); r.text = "The judge is prompted to decompose each answer into atomic claims and check each one individually against the context — a single unsupported clause inside an otherwise-correct long answer must still be flagged."
r.font.size = Pt(11.5); r.font.name = SANS; r.font.color.rgb = INK_SOFT
add_textbox(s, Inches(0.6), Inches(6.0), Inches(11.5), Inches(0.4),
            "n = 30 calibration rows, held out from the evaluation set. These are framework-validation numbers — not the strategy comparison.",
            size=11, italic=True, color=MUTED, font=SANS)
add_footer(s, 9)


# ===========================================================================
# SLIDE 10 -- Project Timeline
# ===========================================================================
s = add_slide()
add_header(s, "07 · Project Timeline", "Four review checkpoints across the semester")

phases = [
    ("PHASE 1", "Zeroth Review", "Foundation", "done",
     ["Problem identification", "Literature survey", "Initial proposal & scope"]),
    ("PHASE 2", "First Review", "Pipeline & Framework", "current",
     ["Fixed RAG pipeline built", "All 4 strategies implemented", "Evaluation framework validated"]),
    ("PHASE 3", "Second Review", "Full-Scale Benchmark", "upcoming",
     ["Run all strategies at scale", "Final accuracy / hallucination comparison", "Cost & latency analysis"]),
    ("PHASE 4", "Final Review", "Analysis & Submission", "upcoming",
     ["Statistical significance check", "Final report & recommendations", "Project documentation"]),
]
cw, gap, x0, top, h = 2.84, 0.25, 0.6, 2.15, 3.75
status_style = {
    "done": (DONE_GREEN_BG, DONE_GREEN, "✓ DONE"),
    "current": (RGBColor(0xDC, 0xE7, 0xF3), NAVY, "● IN PROGRESS"),
    "upcoming": (GREY_BG, GREY_INK, "○ UPCOMING"),
}
centers = []
for i, (phase_no, review, title, status, items) in enumerate(phases):
    x = x0 + i * (cw + gap)
    centers.append(x + cw / 2)
    fill, ink, pill_label = status_style[status]
    card = add_rect(s, Inches(x), Inches(top), Inches(cw), Inches(h), fill=WHITE, line_color=LINE, radius=0.06)
    tf = card.text_frame; tf.margin_left = Inches(0.16); tf.margin_top = Inches(0.16); tf.margin_right = Inches(0.14); tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = phase_no; r.font.size = Pt(9.5); r.font.bold = True; r.font.name = SANS; r.font.color.rgb = MUTED
    p2 = tf.add_paragraph(); p2.space_before = Pt(4)
    r2 = p2.add_run(); r2.text = review; r2.font.size = Pt(13.5); r2.font.bold = True; r2.font.name = SANS; r2.font.color.rgb = INK
    p3 = tf.add_paragraph(); p3.space_before = Pt(2)
    r3 = p3.add_run(); r3.text = title; r3.font.size = Pt(11.5); r3.font.italic = True; r3.font.name = SANS; r3.font.color.rgb = ink
    for it in items:
        pi = tf.add_paragraph(); pi.space_before = Pt(8)
        ri = pi.add_run(); ri.text = f"—  {it}"; ri.font.size = Pt(10); ri.font.name = SANS; ri.font.color.rgb = INK_SOFT

    pill_w = Inches(0.26 + 0.078 * len(pill_label))
    pill = add_rect(s, Inches(x + 0.14), Inches(top + h - 0.42), pill_w, Inches(0.3), fill=fill, radius=0.5)
    ptf = pill.text_frame; ptf.margin_left = Inches(0.08); ptf.margin_right = Inches(0.08)
    ptf.margin_top = 0; ptf.margin_bottom = 0; ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pp = ptf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    pr = pp.add_run(); pr.text = pill_label; pr.font.size = Pt(9); pr.font.bold = True; pr.font.name = SANS; pr.font.color.rgb = ink

for i in range(3):
    connector(s, centers[i] + cw / 2 - 0.15, top + h / 2, centers[i + 1] - cw / 2 + 0.15, top + h / 2, color=MUTED)
add_footer(s, 10)


# ===========================================================================
# SLIDE 11 -- Tools & Technologies Used
# ===========================================================================
s = add_slide()
add_header(s, "08 · Tools & Technologies", "What it's built with")

stack = [
    ("Orchestration", "LangChain (LCEL)"),
    ("Vector store", "ChromaDB"),
    ("Embeddings", "text-embedding-3-small"),
    ("Base LLM", "gpt-4o-mini"),
    ("Validation", "Pydantic v2"),
    ("Evaluation", "RAGAS"),
    ("Dataset", "RAGTruth (Hugging Face)"),
    ("Backend / GUI", "FastAPI + WebSocket + JS"),
    ("Testing", "pytest (14 tests)"),
]
cols3 = 3
cw2, ch2 = 3.85, 1.3
gx, gy = 0.35, 0.28
x0b, y0b = 0.6, 2.15
for i, (role, name) in enumerate(stack):
    r_, c_ = divmod(i, cols3)
    x = x0b + c_ * (cw2 + gx)
    y = y0b + r_ * (ch2 + gy)
    card = add_rect(s, Inches(x), Inches(y), Inches(cw2), Inches(ch2), fill=RGBColor(0xF7, 0xF9, 0xFA), line_color=LINE, radius=0.1)
    tf = card.text_frame; tf.margin_left = Inches(0.22); tf.margin_top = Inches(0.2); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = role.upper(); r.font.size = Pt(10); r.font.bold = True; r.font.name = SANS; r.font.color.rgb = MUTED
    p2 = tf.add_paragraph(); p2.space_before = Pt(4)
    r2 = p2.add_run(); r2.text = name; r2.font.size = Pt(15.5); r2.font.bold = True; r2.font.name = SANS; r2.font.color.rgb = NAVY
add_footer(s, 11)


prs.save(str(OUT_PATH))
print(f"Saved -> {OUT_PATH}")
